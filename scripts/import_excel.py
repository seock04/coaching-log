#!/usr/bin/env python3
"""Import the KSC coaching log workbook without copying personal data into Git."""

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


TIME_RANGE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*[-~]\s*(\d{1,2}):(\d{2})\s*$")


def number(value):
    return int(value or 0)


def parse_workbook(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook["코칭일지"]
    records = []
    errors = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=7, max_row=546, max_col=10, values_only=True), start=7
    ):
        date, time_range, client, paid, free, received, format_name, given, mentor, milestone = row
        if not (date and hasattr(date, "year") and client):
            continue
        match = TIME_RANGE.match(str(time_range or ""))
        if not match:
            errors.append(f"{row_number}행: 시간 형식을 읽을 수 없습니다: {time_range!r}")
            continue
        sh, sm, eh, em = map(int, match.groups())
        records.append(
            {
                "session_date": date.date().isoformat(),
                "start_time": f"{sh:02d}:{sm:02d}:00",
                "end_time": f"{eh:02d}:{em:02d}:00",
                "client_name": str(client).strip(),
                "paid_minutes": number(paid),
                "free_minutes": number(free),
                "received_coach_the_coach_minutes": number(received),
                "coaching_format": str(format_name or "").strip(),
                "given_coach_the_coach_minutes": number(given),
                "mentor_coaching_minutes": number(mentor),
                "milestone": str(milestone or "").strip(),
                "notes": "",
            }
        )
    return records, errors


def post_batch(url, key, user_id, records):
    payload = [{**record, "user_id": user_id} for record in records]
    request = urllib.request.Request(
        f"{url.rstrip('/')}/rest/v1/coaching_sessions?on_conflict=user_id,session_date,start_time,end_time,client_name",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Content-Type": "application/json",
            "Prefer": "resolution=merge-duplicates,return=minimal",
        },
        method="POST",
    )
    with urllib.request.urlopen(request) as response:
        return response.status


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--user-id", required=True, help="Supabase Auth user UUID")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    records, errors = parse_workbook(args.workbook)
    if errors:
        print("\n".join(errors), file=sys.stderr)
        return 2
    totals = {
        key: sum(record[key] for record in records)
        for key in [
            "paid_minutes", "free_minutes", "received_coach_the_coach_minutes",
            "given_coach_the_coach_minutes", "mentor_coaching_minutes",
        ]
    }
    print(json.dumps({"records": len(records), "totals": totals}, ensure_ascii=False, indent=2))
    if args.dry_run:
        return 0
    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        print("SUPABASE_URL과 SUPABASE_SERVICE_ROLE_KEY가 필요합니다.", file=sys.stderr)
        return 2
    try:
        for start in range(0, len(records), 100):
            post_batch(url, key, args.user_id, records[start:start + 100])
    except urllib.error.HTTPError as error:
        print(error.read().decode("utf-8"), file=sys.stderr)
        return 1
    print(f"{len(records)}개 기록을 가져왔습니다.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
